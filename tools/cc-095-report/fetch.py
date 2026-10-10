#!/usr/bin/env python3
"""CC-095: download the data for the check and write extracts to data/cc-095/. Network needed.

Eurostat (dissemination API, JSON):
  prc_hicp_minr  HICP, ECOICOP ver.2, monthly indices (2025=100) and annual rates of change. prc_hicp_manr, named in
                 older notes, stops at 2025-12 (Eurostat moved the HICP to ECOICOP ver.2 in 2026).
                 -> eurostat_hicp_energy_eu_2025_2026.csv (every EU country, Jan 2025 to the latest month)
                 -> eurostat_hicp_energy_mt_eu_history.csv (Malta and the EU, Jan 2014 to the latest month)
  nrg_pc_204     household electricity prices, band DC (2,500-4,999 kWh), all taxes, EUR per kWh, 2019-S1 on
                 -> eurostat_nrg_pc_204_dc.csv
  nrg_pc_205     non-household electricity prices, band ID (500-1,999 MWh), excluding VAT, EUR per kWh
                 -> eurostat_nrg_pc_205_id.csv
  gov_10a_exp    COFOG 04.3 "Fuel and energy": subsidies (D.3) and total expenditure, EUR million and % of GDP
                 -> eurostat_cofog_fuel_energy.csv
  nama_10_gdp    GDP at current prices, Malta -> eurostat_nama_10_gdp_mt.csv
European Commission, Weekly Oil Bulletin (prices with and without taxes, history workbook):
                 -> oil_bulletin_mt_it_eu_weekly.csv (Malta, Italy, EU weighted average; 2019 on)
                 -> oil_bulletin_latest_all.csv (every country, the latest week, with taxes)
Hashes of the downloaded files go to raw_files_sha256.csv. Raw files are kept in out/raw/ (git-ignored).
"""
import csv
import datetime as dt
import hashlib
import json
import pathlib
import subprocess

HERE = pathlib.Path(__file__).resolve().parent
D = HERE.parents[1] / "data" / "cc-095"
RAW = HERE / "out" / "raw"
RAW.mkdir(parents=True, exist_ok=True)
D.mkdir(parents=True, exist_ok=True)
B = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/"
UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
TODAY = dt.date.today().isoformat()
EU27 = "AT BE BG CY CZ DE DK EE EL ES FI FR HR HU IE IT LT LU LV MT NL PL PT RO SE SI SK".split()
HASHES = []


def curl(url, path):
    subprocess.run(["curl", "-sSL", "--fail", "--retry", "4", "--retry-all-errors", "--retry-delay", "5",
                    "--max-time", "300", "-A", UA, "-o", str(path), url], check=True)
    HASHES.append({"file": path.name, "url": url, "bytes": path.stat().st_size,
                   "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "retrieved": TODAY})
    return path


def eurostat(name, query):
    """Flatten a Eurostat JSON-stat response into rows (dimension order from d['id'])."""
    url = B + query + "&format=JSON&lang=en"
    d = json.loads(curl(url, RAW / f"{name}.json").read_text())
    ids, sizes = d["id"], d["size"]
    cats = {k: {i: c for c, i in d["dimension"][k]["category"]["index"].items()} for k in ids}
    rows = []
    for k, v in d["value"].items():
        kk, pos = int(k), {}
        for dim, n in zip(reversed(ids), reversed(sizes)):
            pos[dim] = cats[dim][kk % n]
            kk //= n
        pos.update(value=v, flag=d.get("status", {}).get(k, ""), dataset=query.split("?")[0],
                   updated=d.get("updated", ""), retrieved=TODAY, url=url)
        rows.append(pos)
    return rows


def write(name, rows, keys):
    with open(D / name, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=keys, lineterminator="\r\n", extrasaction="ignore")
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: tuple(str(r.get(k, "")) for k in keys)))
    print(f"{name}: {len(rows)} rows")


# ---------------------------------------------------------------- HICP (ECOICOP ver.2)
CODES = ["TOTAL", "NRG", "ELC_GAS", "FUEL", "CP0451", "CP0452", "CP04522", "CP0722", "CP07221", "CP07222", "AP_NRG"]
geo = "&".join(f"geo={g}" for g in EU27 + ["EU27_2020", "EA20"])
coi = "&".join(f"coicop18={c}" for c in CODES)
h1 = eurostat("hicp_eu", f"prc_hicp_minr?{geo}&{coi}&unit=I25&unit=RCH_A&sinceTimePeriod=2025-01")
write("eurostat_hicp_energy_eu_2025_2026.csv", h1,
      ["coicop18", "unit", "geo", "time", "value", "flag", "dataset", "updated", "retrieved"])
h2 = eurostat("hicp_mt_eu", f"prc_hicp_minr?geo=MT&geo=EU27_2020&{coi}&unit=I25&unit=RCH_A&sinceTimePeriod=2014-01")
write("eurostat_hicp_energy_mt_eu_history.csv", h2,
      ["coicop18", "unit", "geo", "time", "value", "flag", "dataset", "updated", "retrieved"])

# ---------------------------------------------------------------- electricity prices, households
e = eurostat("nrg_pc_204", "nrg_pc_204?nrg_cons=KWH2500-4999&tax=I_TAX&unit=KWH&currency=EUR&sinceTimePeriod=2019-S1")
write("eurostat_nrg_pc_204_dc.csv", e,
      ["geo", "time", "value", "flag", "nrg_cons", "tax", "currency", "dataset", "updated", "retrieved"])
# non-household (business) electricity, band ID (500-1,999 MWh), excluding VAT and other recoverable taxes
nb = eurostat("nrg_pc_205", "nrg_pc_205?nrg_cons=MWH500-1999&tax=X_VAT&unit=KWH&currency=EUR&sinceTimePeriod=2019-S1")
write("eurostat_nrg_pc_205_id.csv", nb,
      ["geo", "time", "value", "flag", "nrg_cons", "tax", "currency", "dataset", "updated", "retrieved"])

# ---------------------------------------------------------------- COFOG 04.3 fuel and energy; GDP
geo2 = "&".join(f"geo={g}" for g in EU27 + ["EU27_2020"])
c = eurostat("cofog", f"gov_10a_exp?{geo2}&sector=S13&cofog99=GF0403&na_item=D3&na_item=TE&unit=MIO_EUR"
                      f"&unit=PC_GDP&sinceTimePeriod=2015")
write("eurostat_cofog_fuel_energy.csv", c,
      ["na_item", "unit", "geo", "time", "value", "flag", "cofog99", "sector", "dataset", "updated", "retrieved"])
g = eurostat("gdp", "nama_10_gdp?geo=MT&na_item=B1GQ&unit=CP_MEUR&sinceTimePeriod=2019")
write("eurostat_nama_10_gdp_mt.csv", g, ["geo", "time", "value", "flag", "na_item", "unit", "dataset", "updated",
                                         "retrieved"])

# ---------------------------------------------------------------- Weekly Oil Bulletin
import openpyxl  # noqa: E402

WOB = ("https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
       "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx")
wb = openpyxl.load_workbook(curl(WOB, RAW / "wob_history.xlsx"), read_only=True, data_only=True)
out = []
for sheet, basis in (("Prices with taxes", "with taxes"), ("Prices wo taxes", "without taxes")):
    rows = list(wb[sheet].iter_rows(values_only=True))
    hdr = rows[0]
    for r in rows[3:]:
        if not isinstance(r[0], dt.datetime) or r[0].year < 2019:
            continue
        for i, h in enumerate(hdr):
            if not h or str(h) == "CTR":
                continue
            ctry, prod = str(h).split("_")[0], str(h).split("_")[-1]
            if ctry not in ("MT", "IT", "EU") or prod not in ("euro95", "diesel", "LPG", "oil"):
                continue
            if r[i] is None:
                continue
            out.append({"date": r[0].date().isoformat(), "geo": ctry, "product": str(h).split("tax_")[-1],
                        "basis": basis, "eur_per_1000": round(float(r[i]), 2), "source": "Weekly Oil Bulletin history",
                        "retrieved": TODAY})
write("oil_bulletin_mt_it_eu_weekly.csv", out,
      ["geo", "product", "basis", "date", "eur_per_1000", "source", "retrieved"])
# latest week, every country, with taxes (from the same workbook, so one vintage)
rows = list(wb["Prices with taxes"].iter_rows(values_only=True))
hdr, latest = rows[0], rows[3]
lat = []
for i, h in enumerate(hdr):
    if h and str(h) != "CTR" and str(h).split("_")[0] in EU27 + ["EU"] and str(h).endswith(("euro95", "diesel")):
        if latest[i] is not None:
            lat.append({"date": latest[0].date().isoformat(), "geo": str(h).split("_")[0],
                        "product": str(h).split("tax_")[-1], "basis": "with taxes",
                        "eur_per_1000": round(float(latest[i]), 2), "retrieved": TODAY})
write("oil_bulletin_latest_all.csv", lat, ["product", "geo", "date", "basis", "eur_per_1000", "retrieved"])

with open(D / "raw_files_sha256.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["file", "url", "bytes", "sha256", "retrieved"], lineterminator="\r\n")
    w.writeheader()
    w.writerows(HASHES)
print("done", TODAY)

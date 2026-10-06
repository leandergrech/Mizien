#!/usr/bin/env python3
"""CC-094: download the EEA datasets behind the Commission's Malta sentences and extract the rows we use.

Raw files go to tools/cc-094-report/out/raw/ (git-ignored); extracts go to data/cc-094/ with the source URL and the
retrieval date on every row. The EEA publishes these files on a public Nextcloud share ("datashare") whose WebDAV
address is https://sdi.eea.europa.eu/datashare/public.php/webdav/<dataset folder>/<file>, with the share's public
link id as the user name and no password. The folders are listed at https://sdi.eea.europa.eu/datastore/public; the
script reads the public link id from that listing rather than storing it.

1. Greenhouse gas emissions under the Effort Sharing Legislation, 2005-2024 (EEA, reporting year 2025, sheet dated
   4 Nov 2025): eea_t_esd-2022_p_2005-2024_v01_r00/GHG_ESD-ESR_2025.xlsx. Status of each value as the EEA labels it:
   2005-2012 EEA estimates from the 2022 inventory and ETS data; 2013-2020 ESD reviews; 2021-2023 ESR reviews
   (2025 comprehensive review); 2024 approximated inventory ("Proxy").
2. Member States' GHG projections, 2025 submission (EEA, file of 13 Apr 2026):
   eea_t_ghg-emission-projections_p_2025_v01_r00/GHG_projections_2025_EEA_incl_pivots.xlsx. Sheet "Totals" holds the
   2025 submission (WEM, WAM), sheet "2023 data" the 2023 submission. The EEA's quality-checked value is in the
   "Gapfilled" column (the "Reported Value" column is empty for these rows).
3. NECPR Annex I, progress to GHG targets, 2025 (EEA, file of 20 Feb 2026):
   eea_t_progress-ghg-targets_p_2025_v01_r00/NECPR_Annex I_GHG Progress_2025.xlsx, Table_2 (Malta rows only).
   Pitfall: its "Annual emission allocation" rows are in Mt although labelled ktCO2e, so the EEA's own "difference"
   rows for Malta are wrong by a factor of 1,000; we use the AEAs from the Commission's implementing decisions
   (data/cc-094/esr_legal_inputs.csv) instead.
"""
import csv, datetime, pathlib, re, subprocess, urllib.parse
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
D = ROOT / "data" / "cc-094"
RAW = pathlib.Path(__file__).resolve().parent / "out" / "raw"
DAV = "https://sdi.eea.europa.eu/datashare/public.php/webdav/"
TODAY = datetime.date.today().isoformat()
FILES = {
    "esr": "eea_t_esd-2022_p_2005-2024_v01_r00/GHG_ESD-ESR_2025.xlsx",
    "proj": "eea_t_ghg-emission-projections_p_2025_v01_r00/GHG_projections_2025_EEA_incl_pivots.xlsx",
    "necpr": "eea_t_progress-ghg-targets_p_2025_v01_r00/NECPR_Annex I_GHG Progress_2025.xlsx",
}


def share_id():
    """The public link id printed in the EEA datastore's web listing (a public download link, not a credential)."""
    html = subprocess.run(["curl", "-sSL", "https://sdi.eea.europa.eu/datastore/public?path=/" + FILES["esr"].split("/")[0] + "/"],
                          check=True, capture_output=True, text=True).stdout
    return re.search(r"datashare/s/([A-Za-z0-9]+)/download", html).group(1)


def url(rel):
    return DAV + urllib.parse.quote(rel)


def get(rel):
    RAW.mkdir(parents=True, exist_ok=True)
    out = RAW / rel.split("/")[-1].replace(" ", "_")
    if not out.exists():
        subprocess.run(["curl", "-sS", "-f", "-u", share_id() + ":", "-o", str(out), url(rel)], check=True)
    return out


def write(name, header, rows):
    with open(D / name, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(name, len(rows), "rows")


D.mkdir(exist_ok=True)

# 1. ESR emissions, all countries, 2005-2024 (long sheet) plus the EEA's status label per period
wb = openpyxl.load_workbook(get(FILES["esr"]), read_only=True, data_only=True)
STATUS = {range(2005, 2013): "EEA estimate (2022 inventory, ETS data, ETS scope correction)",
          range(2013, 2021): "ESD review", range(2021, 2024): "ESR review (2025 comprehensive review)",
          range(2024, 2025): "approximated inventory (proxy)"}
rows = []
it = wb["GHG_ESD"].iter_rows(values_only=True)
next(it)
for country, year, value, legal, unit, src in it:
    if country is None:
        continue
    y = int(year)
    status = next(s for r, s in STATUS.items() if y in r)
    rows.append([country, y, value, legal, unit, status, "EEA GHG_ESD-ESR_2025.xlsx sheet GHG_ESD", url(FILES["esr"]), TODAY])
# the 'ESD baseyear 2005' column of the tabular sheet (the 2013-2020 Effort Sharing Decision base)
for r in wb["ESD ESR emissions"].iter_rows(values_only=True):
    if r and r[0] == "MT" and r[1] == "Malta":
        rows.append(["Malta", 2005, r[2], "ESD base year 2005", "MtCO2 eq",
                     "base year under the Effort Sharing Decision 2013-2020 (older inventory)",
                     "EEA GHG_ESD-ESR_2025.xlsx sheet 'ESD ESR emissions', column 'ESD baseyear 2005'",
                     url(FILES["esr"]), TODAY])
write("eea_esr_emissions.csv", ["country", "year", "value", "series", "unit", "eea_status", "sheet", "source_url",
                                "retrieved"], rows)

# 2. Projections: Malta ESR rows by sector (2025 submission), all countries' ESR totals for 2023-2030,
#    and Malta's 2023 submission totals
wb = openpyxl.load_workbook(get(FILES["proj"]), read_only=True, data_only=True)
rows = []
it = wb["Totals"].iter_rows(values_only=True)
next(it)
for r in it:
    cc, yr, sub, cat, scen, gas, rep, cal, gap = r[:9]
    if gas != "ESR emissions (ktCO2e)" or yr is None:
        continue
    keep = (cc == "MT" and 2023 <= int(yr) <= 2040) or (cat == "Total excluding LULUCF" and 2023 <= int(yr) <= 2030)
    if keep:
        rows.append(["2025 submission (sheet Totals)", cc, int(yr), sub, cat, scen, gap, "ktCO2e",
                     url(FILES["proj"]), TODAY])
it = wb["2023 data"].iter_rows(values_only=True)
next(it)
for yr, cc, sub, cat, scen, gas, gap in (r[:7] for r in it):
    if cc == "MT" and gas == "ESR" and cat == "Total excluding LULUCF":
        rows.append(["2023 submission (sheet 2023 data)", cc, int(yr), sub, cat, scen, gap, "ktCO2e",
                     url(FILES["proj"]), TODAY])
write("eea_ghg_projections_esr.csv", ["submission", "country", "year", "submission_year", "category", "scenario",
                                      "value_gapfilled", "unit", "source_url", "retrieved"], rows)

# 3. NECPR Table 2, Malta (kept to document the unit pitfall and Malta's own reported 2022-2023 ESR emissions)
wb = openpyxl.load_workbook(get(FILES["necpr"]), read_only=True, data_only=True)
rows = []
it = wb["Table_2"].iter_rows(values_only=True)
next(it)
for grp, el, cc, unit, year, value, src in it:
    if cc == "MT":
        rows.append([el, cc, unit, year, value, src, url(FILES["necpr"]), TODAY])
write("eea_necpr2025_table2_malta.csv", ["reporting_element", "country", "unit_as_labelled", "year", "value",
                                         "eea_source_column", "source_url", "retrieved"], rows)

#!/usr/bin/env python3
"""CC-109: download the data behind "transport is 48% of Malta's effort-sharing emissions, up 45% since 2005".

Writes to data/cc-109/ (CSV, with source URL and retrieval date on every row):
  eurostat_env_air_gge.csv    Eurostat env_air_gge (the EEA's compilation of the national GHG inventories, 2026
                              submission), Malta and EU-27, transport sub-categories and totals, GHG and CO2,
                              2005-2024, with Eurostat's flags.
  eea_esr_emissions.csv       EEA "Greenhouse gas emissions under the Effort Sharing Legislation, 2005-2024"
                              (published 6 Nov 2025; doi:10.2909/f80bebef-447e-4882-b3e9-c4a91cbae4f5): Malta and
                              EU-27 effort-sharing totals, with the EEA's status for each year.
  eea_approx_inventory_2024.csv  EEA "GovReg: Approximated estimates for greenhouse gas emissions, 2024" (published
                              5 Nov 2025; doi:10.2909/cd572015-4d0d-408a-abf8-290ed2057bf7): every Malta row and the
                              EU-27 transport and effort-sharing rows (Total, ETS, non-ETS and ESR scopes).
  eea_govreg_2026v1_mt.csv    (--govreg, 95 MB download) EEA "GovReg: National emissions reported to the UNFCCC and
                              to the EU, 2026 ver. 1.0" (data of 15 Mar 2026, published 17 Apr 2026;
                              doi:10.2909/83ee8f8c-1422-4e3f-af63-ba88146811e5): Malta transport rows and totals,
                              the inventory version available when the Commission drafted its report.

The EEA files sit in public Nextcloud shares on sdi.eea.europa.eu; the share token is read from the dataset page and
the file fetched over the share's public WebDAV address.
"""
import csv, datetime, io, json, pathlib, re, sys, tempfile, urllib.request

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
D = ROOT / "data" / "cc-109"
D.mkdir(exist_ok=True)
TODAY = datetime.date.today().isoformat()
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) Mizien fact-check"}


def get(url, auth=None, timeout=300):
    req = urllib.request.Request(url, headers=dict(UA))
    if auth:
        import base64
        req.add_header("Authorization", "Basic " + base64.b64encode(f"{auth}:".encode()).decode())
    return urllib.request.urlopen(req, timeout=timeout)


def write(name, rows):
    with open(D / name, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(name, len(rows), "rows")


# ---------------------------------------------------------------- 1. Eurostat env_air_gge
CRF = ["TOTX4_MEMO", "CRF1A1", "CRF1A2", "CRF1A3", "CRF1A3A", "CRF1A3B", "CRF1A3B1", "CRF1A3B2", "CRF1A3B3",
       "CRF1A3B4", "CRF1A3C", "CRF1A3D", "CRF1A3E", "CRF1A4", "CRF2", "CRF2F", "CRF3", "CRF5"]
Q = ("https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/env_air_gge?format=JSON&lang=en"
     "&geo=MT&geo=EU27_2020&unit=MIO_T&airpol=GHG&airpol=CO2&sinceTimePeriod=2005" + "".join("&src_crf=" + c for c in CRF))


def eurostat():
    d = json.load(get(Q, timeout=120))
    ids, cats = d["id"], []
    for k in ids:
        idx = d["dimension"][k]["category"]["index"]
        cats.append(sorted(idx, key=lambda c: idx[c]))
    lab = d["dimension"]["src_crf"]["category"]["label"]
    rows = []
    for key, v in d["value"].items():
        n, pos = int(key), []
        for s in reversed(d["size"]):
            pos.append(n % s)
            n //= s
        r = {k: cats[i][p] for i, (k, p) in enumerate(zip(ids, reversed(pos)))}
        rows.append({"dataset": "env_air_gge", "geo": r["geo"], "airpol": r["airpol"], "src_crf": r["src_crf"],
                     "label": lab[r["src_crf"]], "year": r["time"], "value": v, "unit": "Mt CO2e",
                     "flag": d.get("status", {}).get(key, ""), "eurostat_updated": d.get("updated"),
                     "retrieved": TODAY, "query_url": Q})
    rows.sort(key=lambda r: (r["geo"], r["airpol"], CRF.index(r["src_crf"]), r["year"]))
    write("eurostat_env_air_gge.csv", rows)


# ---------------------------------------------------------------- EEA shares
def eea_file(uuid, filename):
    page = get(f"https://sdi.eea.europa.eu/data/{uuid}").read().decode("utf-8", "ignore")
    token = re.search(r'sharingToken" value="([^"]+)"', page).group(1)
    url = f"https://sdi.eea.europa.eu/datashare/public.php/webdav/{urllib.request.quote(filename)}"
    return get(url, auth=token), f"https://sdi.eea.europa.eu/data/{uuid} ({filename})"


ESR_UUID = "f80bebef-447e-4882-b3e9-c4a91cbae4f5"
PROXY_UUID = "cd572015-4d0d-408a-abf8-290ed2057bf7"
GOVREG_UUID = "83ee8f8c-1422-4e3f-af63-ba88146811e5"


def esr_status(y):
    # the workbook's own comment: "EEA estimates for 2005-2012; 2013-2020 reviews under ESD; ESR 2021-2023 reviews
    # under ESR, approximated inventory 2024"
    return ("EEA estimate (2005-2012)" if y <= 2012 else "ESD review (2013-2020)" if y <= 2020 else
            "ESR review (2021-2023)" if y <= 2023 else "approximated inventory (proxy)")


def eea_esr():
    f, src = eea_file(ESR_UUID, "GHG_ESD-ESR_2025.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
    rows = []
    for r in wb["GHG_ESD"].iter_rows(min_row=2, values_only=True):
        if r[0] in ("Malta", "EU-27"):
            rows.append({"dataset": "EEA ESD/ESR emissions 2005-2024", "country": r[0], "year": r[1], "value": r[2],
                         "unit": r[4], "regime": r[3], "status": esr_status(int(r[1])), "source": src,
                         "doi": "10.2909/" + ESR_UUID, "published": "2025-11-06", "retrieved": TODAY})
    ws = wb["ESD ESR emissions"]
    for r in ws.iter_rows(min_row=13, values_only=True):   # 2005 base year of the Effort Sharing Decision (column C)
        if r[1] in ("Malta", "EU-27") and r[2] is not None:
            rows.append({"dataset": "EEA ESD/ESR emissions 2005-2024", "country": r[1], "year": "ESD base year 2005",
                         "value": r[2], "unit": "MtCO2 eq", "regime": "ESD", "status": "ESD base year (AR4 GWPs)",
                         "source": src, "doi": "10.2909/" + ESR_UUID, "published": "2025-11-06", "retrieved": TODAY})
    write("eea_esr_emissions.csv", rows)


def eea_proxy():
    f, src = eea_file(PROXY_UUID, "EU_2024_GHG_Approximated_Inventory.xlsx")
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
        tmp.write(f.read())
        tmp.flush()
        wb = openpyxl.load_workbook(tmp.name, read_only=True, data_only=True)
        it = wb["EEA proxy dataset (plus)"].iter_rows(values_only=True)
        hdr = next(it)
        rows, blank = [], 0
        for r in it:
            if r[1] is None:
                blank += 1
                if blank > 1000:
                    break
                continue
            blank = 0
            keep = r[1] == "MT" or (r[1] == "EU27" and r[5] in ("1A3", "Total", "Total_net")
                                    and r[4] in ("Total", "ETS", "non-ETS", "ESR"))
            if keep:
                rows.append({"dataset": "EEA approximated GHG inventory 2024", "country": r[1], "year": r[3],
                             "scope": r[4], "crf_code": r[5], "description": r[6], "sector_name": r[8],
                             "emissions_ms_kt": r[12], "emissions_eea_kt": r[13], "emissions_kt": r[15],
                             "notation": r[16], "provider": r[17], "unit": "kt CO2e", "source": src,
                             "doi": "10.2909/" + PROXY_UUID, "published": "2025-11-05", "retrieved": TODAY})
    write("eea_approx_inventory_2024.csv", rows)


def eea_govreg():
    f, src = eea_file(GOVREG_UUID, "UNFCCC_v30.csv")
    keep_codes = re.compile(r"^(1\.A\.3(\.[a-e])?|1\.A\.1|Sectors/Totals_excl)$")
    rows = []
    for line in io.TextIOWrapper(f, encoding="utf-8-sig"):
        if not line.startswith("MT,"):
            continue
        r = next(csv.reader([line]))
        if keep_codes.match(r[4]) and r[3] in ("All greenhouse gases - (CO2 equivalent)", "CO2") and "2005" <= r[8] <= "2024":
            rows.append({"dataset": "EEA GovReg national emissions 2026 v1.0", "country": r[0], "pollutant": r[3],
                         "sector_code": r[4], "sector_name": r[5], "year": r[8], "value": r[9], "unit": r[7],
                         "notation": r[10], "data_date": r[11], "source": src, "doi": "10.2909/" + GOVREG_UUID,
                         "published": "2026-04-17", "retrieved": TODAY})
    rows.sort(key=lambda r: (r["pollutant"], r["sector_code"], r["year"]))
    write("eea_govreg_2026v1_mt.csv", rows)


if __name__ == "__main__":
    eurostat()
    eea_esr()
    eea_proxy()
    if "--govreg" in sys.argv:
        eea_govreg()

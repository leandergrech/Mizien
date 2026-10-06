#!/usr/bin/env python3
"""CC-113: download the evidence into data/cc-113/ (raw files go to out/raw/, git-ignored).

Sources (all retrieved 6 Oct 2026; re-running records today's date):
  1. EEA indicator "Fossil fuel subsidies in Europe" (published 29 Jan 2025): page state (wording, dates) and its three
     charts: the data packages behind Figure 1 (by energy vector) and Figure 2 (by Member State, 2015 and 2023), and
     the "additional figure" (share of GDP, 2023), whose numbers sit in the chart page itself.
  2. EEA "Europe's environment 2025" country page for Malta, fossil fuel subsidies (published 29 Sep 2025).
  3. European Commission (DG ENER, by Enerdata and Trinomics), "Inventory of energy subsidies in the EU27 - 2024
     edition", the subsidy database published with the 2024 Report on Energy Subsidies in the EU (COM(2025) 17) on
     CIRCABC. This is the dataset the EEA cites ("DG ENER study on energy subsidies").
  4. Eurostat nama_10_gdp (GDP at current prices, EUR million) and gov_10a_main (general government subsidies, Malta).
  5. IMF Fossil Fuel Subsidies Data: 2023 Update (Black et al.), via the IMF Climate Data dashboard (ArcGIS Hub CSV).
  6. OECD/IISD Fossil Fuel Subsidy Tracker, country data (2024 update, file of Jan 2026).
  7. Eurostat nrg_cb_e (electricity supply, Malta 2023): production and imports.
Every CSV keeps the source URL and retrieval date; Eurostat CSVs keep Eurostat's status flags.

Licence: the Commission workbook carries the notice "(c) Copyright Enerdata. Reproduction and diffusion prohibited
(web, photocopy, intranet...) without written permission." So only country-level aggregates (country x year totals,
the GDP series it used) and a handful of Malta figures stated in the report are written to data/cc-113/. The
measure-level extract goes to tools/cc-113-report/out/ (git-ignored) for calc.py's rebuild check. The workbook is
checked against the SHA-256 recorded on 6 Oct 2026 (file modified 9 Apr 2025).
"""
import csv, datetime, hashlib, io, json, pathlib, re, urllib.request, zipfile

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
D = ROOT / "data" / "cc-113"
RAW = pathlib.Path(__file__).resolve().parent / "out" / "raw"
D.mkdir(parents=True, exist_ok=True)
RAW.mkdir(parents=True, exist_ok=True)
TODAY = datetime.date.today().isoformat()
UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"

EEA = "https://www.eea.europa.eu/en/analysis/indicators/fossil-fuel-subsidies"
EEA_SHARE = EEA + "/fossil-fuel-subsidies-as"
EEA_FIG1 = EEA + "/fossil-fuel-subsidies-by-energy"
EEA_FIG2 = EEA + "/fossil-fuel-subsidies-in-eu"
EEA_MT = "https://www.eea.europa.eu/en/europe-environment-2025/countries/malta/fossil-fuel-subsidies"
CIRCABC = "https://circabc.europa.eu/rest/download/4e6f2b2c-702b-429b-bb19-7aeb48f28f58"
CIRCABC_PAGE = ("https://circabc.europa.eu/ui/group/8f5f9424-a7ef-4dbf-b914-1af1d12ff5d2/library/"
                "4e6f2b2c-702b-429b-bb19-7aeb48f28f58/details")
ESTAT = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/{}?format=JSON&lang=en{}"
IMF_CSV = "https://opendata.arcgis.com/datasets/d48cfd2124954fb0900cef95f2db2724_0.csv"
IMF_PAGE = "https://climatedata.imf.org/datasets/d48cfd2124954fb0900cef95f2db2724_0"
TRACKER = ("https://fossilfuelsubsidytracker.org/wp-content/uploads/2026/01/"
           "FossilFuelSubsidiesTracker_CountryData_2024_Update_Jan26.xlsx")
EU27 = "AT BE BG HR CY CZ DK EE FI FR DE EL HU IE IT LV LT LU MT NL PL PT RO SK SI ES SE".split()
NAME2GEO = {"Austria": "AT", "Belgium": "BE", "Bulgaria": "BG", "Croatia": "HR", "Cyprus": "CY", "Czechia": "CZ",
            "Denmark": "DK", "Estonia": "EE", "Finland": "FI", "France": "FR", "Germany": "DE", "Greece": "EL",
            "Hungary": "HU", "Ireland": "IE", "Italy": "IT", "Latvia": "LV", "Lithuania": "LT", "Luxembourg": "LU",
            "Malta": "MT", "Netherlands": "NL", "Poland": "PL", "Portugal": "PT", "Romania": "RO", "Slovakia": "SK",
            "Slovenia": "SI", "Spain": "ES", "Sweden": "SE", "<b>EU-27</b>": "EU27", "EU-27": "EU27"}
ISO2GEO = {"GR": "EL"}   # the Commission inventory and the IMF use GR for Greece; Eurostat uses EL
HASHES = []
OUTDIR = pathlib.Path(__file__).resolve().parent / "out"          # git-ignored
INV_SHA256 = "bf2b412b8edc68dea0bdec6e300c62ba415a2c70cf8fbfd12b6f1431baa88044"   # recorded 6 Oct 2026


def get(url, name):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    data = urllib.request.urlopen(req, timeout=180).read()
    (RAW / name).write_bytes(data)
    HASHES.append({"file": name, "url": url, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(),
                   "retrieved": TODAY})
    return data


def page_state(html):
    """The Plone/Volto page state embedded in an EEA page (window.__data)."""
    i = html.index("window.__data=") + len("window.__data=")
    s = html[i:html.index("</script>", i)].strip().rstrip(";")
    return json.loads(re.sub(r":\s*undefined", ": null", s))["content"]["data"]


def write(name, rows, fields=None):
    with open(D / name, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields or list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(f"{name}: {len(rows)} rows")


# ---------------------------------------------------------------- 1. EEA indicator page and charts
meta = []
for url, name in ((EEA, "eea_page.html"), (EEA_SHARE, "eea_share.html"), (EEA_FIG1, "eea_fig1.html"),
                  (EEA_FIG2, "eea_fig2.html"), (EEA_MT, "eea_malta_soer2025.html")):
    c = page_state(get(url, name).decode("utf-8"))
    meta.append({"page": c.get("title"), "url": url, "effective (published)": c.get("effective"),
                 "modified": c.get("modified"), "created": c.get("created"),
                 "data provenance": "; ".join(f"{p.get('organisation')}: {p.get('title')}"
                                              for p in (c.get("data_provenance") or {}).get("data", [])),
                 "figure note": " ".join(ch.get("text", "") for b in (c.get("figure_note") or [])
                                          for ch in b.get("children", []) if isinstance(ch, dict)).strip(),
                 "retrieved": TODAY})
    if name == "eea_share.html":
        tr = c["visualization"]["data"][0]
        write("eea_share_of_gdp_2023.csv",
              [{"geo": NAME2GEO[n], "country": n.replace("<b>", "").replace("</b>", ""),
                "ffs_share_of_gdp_pct": float(v.rstrip("%")), "year": 2023,
                "source": "EEA, Fossil fuel subsidies as a share of GDP 2023 (chart data)", "url": url,
                "retrieved": TODAY} for n, v in zip(tr["y"], tr["x"])])
    if name == "eea_malta_soer2025.html":
        rows = []
        for tr in c["visualization"]["data"]:
            if not (tr.get("x") and tr.get("y")):   # legend-only or annotation traces carry no data
                continue
            for y, v in zip(tr["x"], tr["y"]):
                if v is not None:
                    rows.append({"series": tr["name"], "year": y, "ffs_share_of_gdp_pct": round(v, 4),
                                 "source": "EEA, Europe's environment 2025, Malta: fossil fuel subsidies (chart data)",
                                 "url": url, "retrieved": TODAY})
        write("eea_malta_trend_soer2025.csv", rows)
write("eea_pages.csv", meta)

for url, name, out, unit_note in ((EEA_FIG2, "eea_fig2_pkg.zip", "eea_fig2_ffs_by_country.csv",
                                   "EUR billion, 2023 prices"),
                                  (EEA_FIG1, "eea_fig1_pkg.zip", "eea_fig1_ffs_by_vector.csv",
                                   "EUR billion, 2023 prices")):
    z = zipfile.ZipFile(io.BytesIO(get(url + "/@@download/file", name)))
    xl = [n for n in z.namelist() if n.endswith("-Data.xlsx")][0]
    ws = openpyxl.load_workbook(io.BytesIO(z.read(xl)), data_only=True)["Original Data"]
    rows = []
    vals = [[v for v in r] for r in ws.iter_rows(values_only=True)]
    hdr = next(r for r in vals if any(isinstance(v, int) and v >= 2015 for v in r if v))
    years = [(j, v) for j, v in enumerate(hdr) if isinstance(v, int)]
    for r in vals[vals.index(hdr) + 1:]:
        lab = r[1]
        if not isinstance(lab, str) or lab.startswith("Note"):
            continue
        for j, y in years:
            if isinstance(r[j], (int, float)):
                rows.append({"label": lab, "geo": NAME2GEO.get(lab, ""), "year": y, "value": r[j], "unit": unit_note,
                             "source": f"EEA chart data package ({xl.split('/')[-1]})", "url": url,
                             "retrieved": TODAY})
    write(out, rows)

# ---------------------------------------------------------------- 3. Commission subsidy inventory (2024 edition)
inv = get(CIRCABC, "ec_energy_subsidies_inventory_2024.xlsx")
if hashlib.sha256(inv).hexdigest() != INV_SHA256:
    print("WARNING: the CIRCABC workbook has changed since 6 Oct 2026 (SHA-256 differs); check the figures again")
wb = openpyxl.load_workbook(io.BytesIO(inv), read_only=True, data_only=True)
SRC_INV = "European Commission (DG ENER; Enerdata, Trinomics), Inventory of energy subsidies in the EU27, 2024 edition"
db = list(wb["Database"].iter_rows(values_only=True))
hdr = db[6]
ix = {h: i for i, h in enumerate(hdr)}
YRS = [y for y in range(2015, 2025) if y in ix]
measures = []
for r in db[7:]:
    if not r[ix["Country Code"]]:
        continue
    row = {"geo": ISO2GEO.get(r[ix["Country Code"]], r[ix["Country Code"]]), "measure": r[ix["Name of policy (English)"]],
           "double_counting": r[ix["Double counting"]], "tbc": r[ix["TBC"]] or "",
           "crisis_measure": r[ix["Created or modified to address energy price rising"]] or "",
           "techno_group": r[ix["Techno group"]], "carrier": r[ix["Main product and carrier"]],
           "sub_carrier": r[ix["Sub-product and sub-carrier"]], "category": r[ix["Subsidy category"]],
           "instrument": r[ix["Subsidy instrument"]], "purpose": r[ix["Purpose"]],
           "sector": r[ix["Main economic sector"]], "env_harm": r[ix["EHES"]],
           "start": str(r[ix["Start of implementation"]] or ""), "end": str(r[ix["End of implementation"]] or ""),
           "actual_costs": r[ix["Actual costs"]], "estimate_costs": r[ix["Estimate costs"]]}
    row.update({f"y{y}": r[ix[y]] if isinstance(r[ix[y]], (int, float)) else "" for y in YRS})
    measures.append(row)
# measure-level rows stay local (Enerdata copyright): every fossil-fuel measure, for calc.py's rebuild check
ffs_rows = [m for m in measures if m["techno_group"] == "Fossil fuels" and m["double_counting"] == "No"]
with open(OUTDIR / "ec_inventory_measures.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(ffs_rows[0]))
    w.writeheader()
    w.writerows(ffs_rows)
print(f"out/ec_inventory_measures.csv (local only, not committed): {len(ffs_rows)} rows")

# the handful of Malta figures the report states (aggregated where the report aggregates)
mt_rows = [m for m in ffs_rows if m["geo"] == "MT"]
facts = []
for label, names, note in (
        ("Energy Support Measures (compensation to Enemalta)", ["Energy Support Measures"],
         "natural gas; income or price support (consumer price guarantee, cost support); created for the price crisis"),
        ("Gas Stabilisation Fund", ["Gas Stabilisation Fund"],
         "natural gas; income or price support; created for the price crisis"),
        ("Cut in excise duty on petrol and diesel", ["Reduction of excise duties on petrol and diesel"],
         "oil; tax expenditure; created for the price crisis; 2023 'to be confirmed' (2022 value used as stand-in)"),
        ("Excise exemptions: inland navigation, fishing, domestic flights (three measures)",
         ["Excise tax exemption on petroleum products consumed in inland water navigation",
          "Excise tax exemption on petroleum products consumed for fishing purpose",
          "Excise tax exemption on kerosene consumed in domestic air traffic"], "oil; tax expenditures"),
        ("Melita TransGas pipeline (Malta-Sicily)",
         ["European Gas Network/Distribution (Interconnexion project btw Malta and Italy)"], "natural gas; grant")):
    sel = [m for m in mt_rows if m["measure"] in names]
    assert len(sel) == len(names), names
    facts.append({"item": label, "y2022": round(sum(float(m["y2022"] or 0) for m in sel), 2),
                  "y2023": round(sum(float(m["y2023"] or 0) for m in sel), 2),
                  "to_be_confirmed": "yes" if any(m["tbc"] == "Yes" for m in sel) else "no",
                  "cost_basis": ", ".join(sorted({"actual costs" if m["actual_costs"] == 1 else "estimate"
                                                  for m in sel})),
                  "classification": note, "unit": "EUR million, 2023 prices",
                  "source": SRC_INV + " (aggregated by Miżien; measure rows not redistributed)", "url": CIRCABC,
                  "retrieved": TODAY})
write("ec_inventory_malta_facts.csv", facts)

# the inventory's own pivot: FFS by country 2018-2023 and "2023 incl. TBC" (FFS End-dates sheet, rows 90-117)
fe = list(wb["FFS End-dates"].iter_rows(values_only=True))
h = next(i for i, r in enumerate(fe) if r[9] == "Row Labels" and r[10] == "Sum of 2018")
piv = []
for r in fe[h + 1:]:
    if r[9] in (None, "Grand Total"):
        break
    for k, col in (("2018", 10), ("2019", 11), ("2020", 12), ("2021", 13), ("2022", 14), ("2023", 15),
                   ("2023 incl. TBC", 16)):
        piv.append({"geo": ISO2GEO.get(r[9], r[9]), "column": k, "ffs_eur_million": r[col],
                    "source": SRC_INV + ", sheet 'FFS End-dates' (FFS by country pivot)", "url": CIRCABC_PAGE,
                    "retrieved": TODAY})
write("ec_inventory_ffs_by_country.csv", piv)

# the inventory's own share-of-GDP pivot 2015-2023 (FFS End-dates sheet, rows 10-37) and the GDP it used ('codes')
h = next(i for i, r in enumerate(fe) if r[0] == "Row Labels" and r[1] == "Sum of 2015")
gpiv = []
for r in fe[h + 1:]:
    if r[0] in (None, "Grand Total"):
        break
    for k, y in enumerate(range(2015, 2024)):
        gpiv.append({"geo": NAME2GEO[r[0]], "year": y, "ffs_eur_million": r[1 + k], "ffs_share_of_gdp": r[11 + k],
                     "source": SRC_INV + ", sheet 'FFS End-dates' (TREND FFS %GDP 2015-2023)", "url": CIRCABC_PAGE,
                     "retrieved": TODAY})
write("ec_inventory_ffs_share_gdp.csv", gpiv)
cd = list(wb["codes"].iter_rows(values_only=True))
g0 = next(i for i, r in enumerate(cd) if r[60] == "GDP")
gh = cd[g0 + 1]
gdp_rows = []
for r in cd[g0 + 2:]:
    if r[60] is None:
        break
    for j, y in enumerate(gh):   # the GDP table sits in columns 60-79 (other tables share these rows)
        if 63 <= j < 79 and isinstance(y, int) and y >= 2015 and isinstance(r[j], (int, float)):
            gdp_rows.append({"geo": ISO2GEO.get(r[61], r[61]) if r[61] != "EU" else "EU27", "year": y,
                             "gdp_eur_million": r[j], "unit": r[62], "source_wording": r[79],
                             "source": SRC_INV + ", sheet 'codes' (GDP)", "url": CIRCABC_PAGE, "retrieved": TODAY})
write("ec_inventory_gdp.csv", gdp_rows)

# ---------------------------------------------------------------- 4. Eurostat


def estat(ds, query):
    d = json.loads(get(ESTAT.format(ds, query), f"eurostat_{ds}.json"))
    dims, sz = d["id"], d["size"]
    cats = [[c for c, _ in sorted(d["dimension"][k]["category"]["index"].items(), key=lambda kv: kv[1])] for k in dims]
    status = d.get("status", {})
    rows = []
    for key, v in d["value"].items():
        n = int(key); pos = []
        for s in reversed(sz):
            pos.append(n % s); n //= s
        pos.reverse()
        rows.append({k: cats[i][p] for i, (k, p) in enumerate(zip(dims, pos))}
                    | {"value": v, "flag": status.get(key, ""), "dataset": ds, "updated": d.get("updated"),
                       "retrieved": TODAY})
    return rows


gdp = [r for r in estat("nama_10_gdp", "&unit=CP_MEUR&na_item=B1GQ&sinceTimePeriod=2015")
       if r["geo"] in EU27 + ["EU27_2020"]]
write("eurostat_nama_10_gdp.csv", gdp)
write("eurostat_gov_10a_main_mt_subsidies.csv",
      estat("gov_10a_main", "&geo=MT&sector=S13&na_item=D3PAY&unit=MIO_EUR&unit=PC_GDP&sinceTimePeriod=2015"))

write("eurostat_nrg_cb_e_mt_2023.csv",
      [r for r in estat("nrg_cb_e", "&geo=MT&siec=E7000&unit=GWH&time=2023") if r["nrg_bal"] in ("GEP", "IMP", "EXP")])

# ---------------------------------------------------------------- 5. IMF fossil fuel subsidies data (2023 update)
imf = list(csv.DictReader(io.StringIO(get(IMF_CSV, "imf_ffs.csv").decode("utf-8-sig"))))
keep = ("Explicit Fossil Fuel Subsidies - Total", "Implicit Fossil Fuel Subsidies - Total",
        "Fossil Fuel Subsidies - Total Implicit and Explicit")
rows = []
for r in imf:
    geo = ISO2GEO.get(r["ISO2"], r["ISO2"])
    if geo in EU27 and r["Indicator"] in keep and r["Unit"] == "Percent of GDP":
        for y in range(2015, 2026):
            rows.append({"geo": geo, "indicator": r["Indicator"], "year": y, "pct_of_gdp": r[f"F{y}"],
                         "source": "Black et al. 2023, IMF Fossil Fuel Subsidies Data: 2023 Update (IMF staff "
                                   "estimates; published Aug 2023, so 2023 and later values were estimated "
                                   "before the year ended)" if r["Source"].startswith("Black") else
                                   r["Source"], "url": IMF_PAGE, "retrieved": TODAY})
write("imf_ffs_eu27.csv", rows)

# ---------------------------------------------------------------- 6. OECD/IISD Fossil Fuel Subsidy Tracker
ws = openpyxl.load_workbook(io.BytesIO(get(TRACKER, "ffs_tracker.xlsx")), read_only=True, data_only=True)["full_data"]
it = ws.iter_rows(values_only=True)
th = next(it)
tx = {h: i for i, h in enumerate(th)}
rows, src = [], {}
for r in it:
    if r[tx["ISO"]] == "MLT":
        rows.append({"geo": "MT", "year": r[tx["Year"]], "mechanism": r[tx["Mechanism"]],
                     "beneficiary": r[tx["Beneficiary"]], "fuel": r[tx["Fuel type"]], "data_source": r[tx["Source"]],
                     "usd_million_nominal": r[tx["USD million, nominal"]],
                     "source": "OECD/IISD Fossil Fuel Subsidy Tracker, country data (2024 update, file Jan 2026)",
                     "url": TRACKER, "retrieved": TODAY})
write("oecd_iisd_tracker_malta.csv", rows)

with open(D / "raw_files_sha256.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(HASHES[0]))
    w.writeheader()
    w.writerows(HASHES)
print("raw files and hashes:", len(HASHES))
